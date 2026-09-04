"""P2 文档解析、分块、图片提取与入库链路测试。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from PIL import Image

from app.core.config import ConfigManager
from app.core.document_service import DocumentService
from app.core.embedder import HashEmbedder
from app.core.text_splitter import RecursiveTextSplitter, estimate_tokens
from app.data.database import Database
from app.data.repository import Repository
from app.data.vector_store import VectorStore


class DocumentPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.config = ConfigManager(self.root)
        self.db = Database(self.root / "app.db")
        self.db.initialize()
        self.repo = Repository(self.db.session())
        self.vector_store = VectorStore(self.root / "chroma")

    def tearDown(self) -> None:
        self.vector_store.close()
        self.repo.close()
        self.db.close()
        self.temp_dir.cleanup()

    def _service(self, ocr=None) -> DocumentService:
        return DocumentService(
            config=self.config,
            database=self.db,
            repository=self.repo,
            vector_store=self.vector_store,
            embedder=HashEmbedder(),
            ocr=ocr or (lambda _path: ""),
        )

    def test_text_splitter_keeps_length_boundary(self) -> None:
        text = "机器学习" * 300
        splitter = RecursiveTextSplitter(chunk_size=100, chunk_overlap=10)
        chunks = splitter.split_text(text)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(estimate_tokens(chunk) <= 150 for chunk in chunks))

    def test_import_markdown_and_search(self) -> None:
        kb = self.repo.create_kb("测试知识库")
        source = self.root / "机器学习.md"
        source.write_text(
            "# 机器学习\n反向传播是神经网络训练的重要算法。\n"
            "梯度下降用于更新模型参数。",
            encoding="utf-8",
        )
        service = self._service()
        docs = service.import_files(kb.id, [source])
        self.assertEqual(docs[0].status, "success")
        query_vector = HashEmbedder().embed(["反向传播 算法 参数"])[0]
        results = self.vector_store.query(kb.id, query_vector, top_k=2, threshold=0.001)
        self.assertTrue(results)

    def test_docx_image_extraction_and_ocr(self) -> None:
        kb = self.repo.create_kb("图片知识库")
        image_path = self.root / "sample.png"
        Image.new("RGB", (64, 32), "white").save(image_path)

        source = self.root / "带图文档.docx"
        from docx import Document

        doc = Document()
        doc.add_paragraph("课程资料正文")
        doc.add_picture(str(image_path))
        doc.save(source)

        service = self._service(ocr=lambda _path: "图中包含：课程结构示意图")
        doc_record = service.import_files(kb.id, [source])[0]
        self.assertEqual(doc_record.status, "success")
        media = self.repo.list_media(document_id=doc_record.id)
        self.assertEqual(len(media), 1)
        self.assertIn("课程结构示意图", media[0].ocr_text)
        self.assertEqual(media[0].status, "success")

    def test_pdf_image_extraction(self) -> None:
        kb = self.repo.create_kb("PDF图片库")
        image_path = self.root / "chart.png"
        Image.new("RGB", (80, 40), "lightblue").save(image_path)
        source = self.root / "带图课件.pdf"

        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((50, 50), "图1 课程结构")
        page.insert_image(pymupdf.Rect(50, 80, 200, 140), filename=str(image_path))
        pdf.save(source)
        pdf.close()

        service = self._service(ocr=lambda _path: "神经网络结构")
        doc_record = service.import_files(kb.id, [source])[0]
        self.assertEqual(doc_record.status, "success")
        media = self.repo.list_media(document_id=doc_record.id)
        self.assertGreaterEqual(len(media), 1)

    def test_direct_image_import_is_indexed(self) -> None:
        kb = self.repo.create_kb("图片资料库")
        source = self.root / "notice.png"
        Image.new("RGB", (100, 40), "white").save(source)
        service = self._service(ocr=lambda _path: "通知：明日上午九点开会")
        doc_record = service.import_files(kb.id, [source])[0]
        self.assertEqual(doc_record.status, "success")
        media = self.repo.list_media(document_id=doc_record.id)
        self.assertEqual(len(media), 1)
        self.assertEqual(media[0].asset_type, "uploaded")
        self.assertIn("明日上午九点开会", media[0].ocr_text)


if __name__ == "__main__":
    unittest.main()
