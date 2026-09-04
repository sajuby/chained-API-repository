"""多格式文档解析器。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


SUPPORTED_TEXT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


@dataclass
class ParsedPage:
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    filename: str
    pages: list[ParsedPage]

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text for page in self.pages).strip()


class DocumentParser:
    def parse(self, file_path: Path) -> ParsedDocument:
        suffix = file_path.suffix.lower()
        if suffix not in SUPPORTED_TEXT_EXTENSIONS:
            raise ValueError(f"不支持的文件格式: {suffix}")
        if suffix == ".pdf":
            return self._parse_pdf(file_path)
        if suffix == ".docx":
            return self._parse_docx(file_path)
        if suffix == ".pptx":
            return self._parse_pptx(file_path)
        return self._parse_plain(file_path)

    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        import pdfplumber

        pages: list[ParsedPage] = []
        with pdfplumber.open(file_path) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                pages.append(ParsedPage(page_number=index, text=text))
        return ParsedDocument(filename=file_path.name, pages=pages)

    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        from docx import Document

        doc = Document(file_path)
        lines: list[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return ParsedDocument(
            filename=file_path.name,
            pages=[ParsedPage(page_number=1, text="\n".join(lines))],
        )

    def _parse_pptx(self, file_path: Path) -> ParsedDocument:
        from pptx import Presentation

        prs = Presentation(file_path)
        pages: list[ParsedPage] = []
        for index, slide in enumerate(prs.slides, start=1):
            lines: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            lines.append(" | ".join(cells))
            pages.append(ParsedPage(page_number=index, text="\n".join(lines)))
        return ParsedDocument(filename=file_path.name, pages=pages)

    def _parse_plain(self, file_path: Path) -> ParsedDocument:
        text = file_path.read_text(encoding="utf-8", errors="replace").strip()
        return ParsedDocument(
            filename=file_path.name,
            pages=[ParsedPage(page_number=1, text=text)],
        )

